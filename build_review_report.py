from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r"C:\Users\32933\Documents\ChatGPT\论文\任务书、开题报告_整体审核意见.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203040"
MUTED = "666666"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_GOLD = "FFF4CE"
PALE_RED = "FCE8E6"
RED = "9B1C1C"
GOLD = "7A5A00"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=11, bold=None, color=INK, latin="Calibri", east_asia="宋体"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def style_callout_paragraph(paragraph, fill=PALE_BLUE, border="D9E2F3"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), border)
        p_bdr.append(el)
    paragraph.paragraph_format.left_indent = Inches(0.10)
    paragraph.paragraph_format.right_indent = Inches(0.10)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)


def add_label_para(doc, label, text, label_color=DARK_BLUE, after=4, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(label)
    set_run_font(r, bold=True, color=label_color)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_rec(doc, title, location, problem, advice, sample=None, priority="P1"):
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run(title)
    if priority == "P0":
        r.font.color.rgb = RGBColor.from_string(RED)
    add_label_para(doc, "位置：", location)
    add_label_para(doc, "问题：", problem, label_color=RED)
    add_label_para(doc, "建议：", advice)
    if sample:
        p2 = doc.add_paragraph()
        style_callout_paragraph(p2, PALE_BLUE)
        r = p2.add_run("可直接参考的改写：\n")
        set_run_font(r, bold=True, color=DARK_BLUE)
        r = p2.add_run(sample)
        set_run_font(r, color=INK)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    set_run_font(run, size=9, color=MUTED)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

# Preset: standard_business_brief. Named override: Chinese text uses SimSun.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("毕业设计材料审阅｜整体审核意见")
set_run_font(hr, size=9, color=MUTED)

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("审核建议  ·  第 ")
set_run_font(fr, size=9, color=MUTED)
add_page_field(fp)
fr = fp.add_run(" 页")
set_run_font(fr, size=9, color=MUTED)

# Title block: memo masthead, without decorative header border.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(5)
r = p.add_run("《任务书、开题报告》")
set_run_font(r, size=24, bold=True, color=DARK_BLUE, east_asia="黑体")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("整体审核意见及针对性修改建议")
set_run_font(r, size=18, bold=True, color=BLUE, east_asia="黑体")

for label, value in (
    ("审核对象：", "AI短剧《山海契灵》创作阐述（任务书与开题报告合并文档）"),
    ("审核范围：", "内容逻辑、两部分一致性、研究与创作方案、进度、参考文献、语言及文档规范"),
    ("建议用途：", "开题答辩前集中修改与自查"),
):
    add_label_para(doc, label, value, after=2)

p = doc.add_paragraph()
style_callout_paragraph(p, PALE_GOLD, border="E6D78A")
r = p.add_run("总体判断：")
set_run_font(r, bold=True, color=GOLD)
r = p.add_run("选题方向具有可行性，传统文化考据、微短剧叙事与AIGC生产三条线也已初步建立；但当前版本存在身份/分工、时间节点、成果形态三类关键矛盾，技术指标和研究方法尚未形成可验收标准，文献综述偏“逐篇介绍”。建议先处理P0问题，再补足论证与格式。")
set_run_font(r)

doc.add_heading("一、优先修改清单", level=1)
rows = [
    ("P0", "任务书“独立完成” vs. 报告“小组/我们”", "明确个人项目还是团队项目；统一全文主语并写清个人独立成果"),
    ("P0", "任务书与报告进度不一致", "统一选题、剧本、素材生成、粗剪、修改、定稿的日期与交付物"),
    ("P0", "成果与技术指标不可验收", "补充集数、单集时长、画幅、分辨率、字幕、声音、脚本和过程材料等指标"),
    ("P0", "文档仍含修订痕迹和异常空段", "提交前接受/拒绝全部修订，删除空段并清理个人元数据"),
    ("P1", "文献综述缺少归纳、比较和研究缺口", "按四类研究线索重写，并说明现有研究未解决什么"),
    ("P1", "AIGC流程、评价方法和合规方案过于笼统", "补充工作流、连续性控制、测试指标、版权与AI标识方案"),
]
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_geometry(t, [900, 3200, 5260])
hdr = t.rows[0].cells
for i, text in enumerate(("级别", "核心问题", "修改目标")):
    hdr[i].text = text
    set_cell_shading(hdr[i], LIGHT)
    hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in hdr[i].paragraphs[0].runs:
        set_run_font(run, bold=True, color=DARK_BLUE)
set_repeat_table_header(t.rows[0])
for level, issue, target in rows:
    cells = t.add_row().cells
    for i, text in enumerate((level, issue, target)):
        cells[i].text = text
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cells[i].paragraphs[0].runs:
            set_run_font(run, color=RED if i == 0 and level == "P0" else INK, bold=(i == 0))
        set_cell_margins(cells[i])

doc.add_heading("二、任务书部分的针对性建议", level=1)
add_rec(
    doc,
    "统一题目与最终成果边界",
    "任务书首页“论文（设计）题目”及“毕业设计的内容”首段。",
    "题目使用“AI短剧”，正文又使用“AIGC短剧”“AI微短剧”，成果有时表述为“20集完整作品”，有时又像“策划方案/创作阐述”。名称和成果边界不统一，会直接影响验收。",
    "先确认最终提交的是20集成片、若干集样片，还是20集策划加部分成片；随后统一题目、任务内容、技术指标、进度和提纲中的术语。若确实提交完整成片，建议统一使用“AIGC微短剧”。",
    "题目可改为：AIGC微短剧《山海契灵》创作设计与阐述。若只完成样片，应明确为：AIGC微短剧《山海契灵》系列策划及样片创作阐述。",
    "P0",
)
add_rec(
    doc,
    "把“主要问题”改成可解决的设计问题",
    "任务书“主要问题”两条。",
    "现有表述是“如何学习技巧”“需要学习知识”，属于学习任务，不是毕业设计要解决的核心问题。",
    "改成三项设计问题，分别对应文化考据、短剧叙事、AI生产；每项问题都要能在设计阐述中给出方法、过程证据和结果。",
    "（1）如何建立《山海经》异兽形象的可追溯考据机制，并完成适度的当代叙事转化；（2）如何用“主线+单元故事”结构兼顾竖屏微短剧节奏与人文表达；（3）如何针对生成式视频的连续性与复杂动作限制，建立角色、场景和镜头的一致性控制流程。",
)
add_rec(
    doc,
    "补齐量化技术指标",
    "任务书“拟达到的要求或技术指标”。",
    "当前五条主要是学术纪律与通用写作要求，无法判断作品是否达标，也没有体现AI短剧的专业特征。",
    "在保留学校通用要求的基础上，增加成片、剧本、视觉、声音、过程材料和设计阐述六类验收指标。数值应以学院要求和实际制作能力为准，不要为了显得完整而虚报。",
    "示例指标：计划完成20集，每集约60—90秒；竖屏9:16，1080×1920；完成20集分集梗概、分镜/镜头清单、统一角色设定表、异兽考据卡、关键提示词迭代记录；字幕无明显错漏，音乐/字体/素材授权可追溯；设计阐述不少于5000字。",
    "P0",
)
add_rec(
    doc,
    "修正计划进度与成果顺序",
    "任务书“计划进度安排”，并与开题报告3.1对照。",
    "开题报告写“2026年7月确定选题”“2026年9—10月确定选题”，任务书却写9月上旬确定；报告又计划2026年12月“最终完成20集完整作品”，但2027年1月才“完成初稿”，先后顺序矛盾。",
    "统一为“前期考据与剧本—视觉设定测试—素材生成—粗剪—内审—二稿—定稿—答辩”的顺序。每一时间段只写可核验交付物，并与任务书学校节点完全对齐。",
    "建议：9月完成选题与开题；10—11月完成考据表、世界观、人物小传、20集分集梗概和剧本初稿；12月完成风格测试及首批样片；1—3月完成素材生成、粗剪和修改；4月完成成片与设计阐述；5月答辩与归档。",
    "P0",
)
add_rec(
    doc,
    "扩充“应收集资料”并形成证据链",
    "任务书“应收集的资料及主要参考文献”。",
    "资料要求只列了两类，缺少政策规范、平台规格、AI工具能力测试、市场样本和版权授权材料。",
    "增加五类资料：权威校注本及异兽原文；AIGC/微短剧相关法规和平台规则；同类作品样本表；AI工具测试记录；音乐、字体、配音、图片等授权证明。形成“来源—创作转化—成片镜头”的对应关系。",
)

doc.add_heading("三、开题报告部分的针对性建议", level=1)
add_rec(
    doc,
    "为题目背景补充证据，减少绝对化判断",
    "1.1“题目背景”两段。",
    "“受到大量观众喜爱”“市场上部分作品存在魔改”“成为重中之重”等判断缺少数据或文献支撑，语气偏宣传化。",
    "至少补充一组行业数据或权威报告、一组平台/作品样本，并说明样本范围与观察维度。把价值判断改为可论证表达。",
    "可改为：基于对××平台在××时间段内×部国风志怪微短剧的观察，部分作品存在异兽设定与原典出处脱节、叙事依赖视觉奇观等现象。因此，本项目拟以可追溯考据和单元叙事为切入点，探索传统古籍IP的AIGC微短剧转化路径。",
)
add_rec(
    doc,
    "降低“完全还原原著”的承诺强度",
    "1.2“选题意义”第一点。",
    "“所有异兽的形貌、习性、出场背景均能对应古籍原文记载”承诺过强。《山海经》版本、注本和图像传统复杂，影视创作也必然包含转化。",
    "将“完全对应”改为“核心特征可追溯”，并建立异兽考据卡：原文、篇目/版本、核心特征、改编内容、改编理由、对应集数。这样既守住依据，也保留创作空间。",
    "核心形貌、能力与文化寓意以权威校注本为依据；新增关系、场景和情节作为创作性转化单独标注，并解释其叙事功能。",
)
add_rec(
    doc,
    "将文献综述从“逐篇摘要”改为“研究脉络+缺口”",
    "1.3.1“论文著作研究成果梳理”。",
    "现有内容主要按阅读顺序介绍单篇文献，缺少观点比较、争议梳理和项目研究缺口；“共42篇”的筛选过程也不可复核。",
    "按四条线索重组：①《山海经》文本考据与图像谱系；②古籍IP影视化与跨媒介转化；③微短剧/单元叙事与竖屏节奏；④AIGC影像生产、美学与风险。每组最后写“共识—分歧—不足—本项目回应”。补充数据库、关键词、检索日期、初筛与纳入数量。",
    "研究缺口示例：现有研究较多讨论传统文化的AI影像转化价值，较少把异兽考据、单元叙事节奏与生成式视频连续性控制放在同一创作流程中验证；本项目将通过考据卡、分集结构表和提示词迭代记录形成实践证据。",
    "P0",
)
add_rec(
    doc,
    "修正影视作品梳理的结构与样本说明",
    "1.3.2“影视作品梳理”。",
    "开头写“主要分为三个部分”，正文实际列出第一至第四，共四部分；同时声称调研27部志怪微短剧、11部AI短剧，却没有作品清单、平台、时间范围和评价维度。",
    "把“三个部分”改为“四个部分”；新增附表或附录，至少列出作品名、平台/年份、题材、单集时长、画幅、钩子、转折、悬念、AI表现问题、可借鉴点。正文只总结规律，不堆作品名。",
    None,
    "P0",
)
add_rec(
    doc,
    "让故事方案具体到可评审的创作单位",
    "2.1“主要内容”。",
    "目前仍是通用的“灭门—查真相—打反派—开放结局”框架，无法看出《山海经》考据如何进入剧情，也无法验证20集是否能支撑主线。",
    "补充一句话梗概、主题命题、主角欲望与缺陷、反派动机、主角弧光；再制作20集分集矩阵，列出每集异兽、原典出处、人性议题、主线推进、开篇钩子和结尾悬念。",
    "一句话梗概模板：为查清宗门覆灭真相，____（主角身份/缺陷）必须在____（期限/限制）内与被误解的异兽结契；每破解一桩“人祸伪装成兽祸”的事件，便更接近____（核心真相），也被迫改变其____（原有偏见）。",
)
add_rec(
    doc,
    "补充明确的研究/设计方法",
    "2.2“思路及方案”，建议新增“研究与创作方法”小节。",
    "目前只有“查资料—构思—AI创作—后期剪辑”的过程描述，没有说明如何分析、比较和评价，因此设计阐述的学术性不足。",
    "建议写明：文本考据法、案例分析法、比较研究法、创作实践法和小规模受众测试法。为每种方法说明材料、步骤、产出和评价指标。",
    "例如：选取不少于×部作品开展编码分析；用5—8名目标受众进行内部试看，从剧情理解、节奏、角色一致性、恐怖氛围和文化信息识别五个维度评分，并记录开放式反馈。样本量不宜夸大，重点是流程透明。",
)
add_rec(
    doc,
    "细化AIGC生产流程与连续性控制",
    "2.2“AI创作：主要选择豆包与即梦等AI工具进行创作”。",
    "仅列工具名称不能说明技术路线；也未说明角色一致性、异兽形态稳定、镜头衔接、失败素材筛选和版本记录方法。",
    "补成完整工作流：剧本拆镜—视觉圣经—角色/异兽设定图—关键帧—图生视频—镜头筛选—补帧/剪辑—配音配乐—字幕—输出。记录工具版本、关键参数、种子/参考图、提示词迭代和淘汰原因。",
    "连续性指标可包括：同角色服饰/面部核心特征一致；同场景色温与构图基调一致；异兽关键形态特征在连续镜头中保留；对明显肢体畸变、穿模、口型错位设定淘汰规则。",
)
add_rec(
    doc,
    "处理个人分工与独立完成要求的冲突",
    "2.2中“我们”“小组”“我的主要工作”，并对照任务书“学生本人独立完成”。",
    "这是当前最需要澄清的合规问题。若是团队作品，任务书的“独立完成”应解释为个人独立承担可考核模块；若是个人作品，则全文不应出现小组主语。",
    "向指导教师确认培养方案允许的协作方式。团队项目应增加成员分工表、共同成果与个人成果边界、素材交接方式，并在阐述中重点证明本人负责模块；个人项目则统一改为“本人/本项目”。",
    "示例：本项目为团队协作作品。本人独立负责AI视觉生成、镜头筛选与后期剪辑，并提交提示词迭代记录、镜头版本表和工程文件作为个人工作证明；剧本/声音等协作内容按实际成员及工作量另行列明。",
    "P0",
)
add_rec(
    doc,
    "增加版权、伦理和平台合规方案",
    "2.2“AI创作”“后期剪辑”之后，建议新增“合规与风险控制”。",
    "缺少对原著版本、参考图、字体、音乐、配音、人物肖像、AI生成内容标识和工具服务条款的说明。",
    "建立素材授权台账；不直接模仿在世艺术家或具体作品风格，不使用未授权真人肖像/声音；核对工具商用条款；按学校与发布平台要求标识AI生成/辅助生成内容；保留生成记录和工程源文件。",
)
add_rec(
    doc,
    "重写工作进度，避免“成片先于初稿”",
    "3.1“工作进度”。",
    "2026年12月写“最终完成20集完整微短剧作品”，2027年1月又“完成初稿”，逻辑倒置；2—4月仅写“二稿/三稿/定稿”，没有对应交付物。",
    "把每月任务拆成“内容成果+技术成果+审核动作”。例如：12月完成视觉设定与3—5集样片，不宜直接承诺全部成片；1月完成20集粗剪；2月根据内审完成二稿；3月完成受众测试与三稿；4月完成成片、阐述与归档材料。",
    None,
    "P0",
)
add_rec(
    doc,
    "优化设计阐述提纲",
    "3.2标题及目录。",
    "标题写作“工作阐述”，与上位标题“设计阐述提纲”不一致；第2章偏过程罗列，第3章只有感悟，缺少方法、效果评价、合规与局限。",
    "统一为“3.2设计阐述提纲”，将论文结构调整为：绪论—创作依据与方法—方案设计—AIGC生产与后期—作品评价—反思与结论。",
    "建议目录：1 绪论（背景、意义、综述、方法）；2 创作依据（异兽考据、案例分析、设计原则）；3 《山海契灵》方案设计（世界观、人物、20集结构、视听风格）；4 AIGC生产与后期（工作流、一致性控制、声音字幕、合规）；5 作品测试与效果分析；6 创作反思与局限；结语、参考文献、附录、致谢。",
)

doc.add_heading("四、参考文献与学术规范建议", level=1)
add_rec(
    doc,
    "补充一手文献并核对二手转引",
    "任务书参考文献[13][16]及1.3.1相关论述。",
    "正文实际使用《古本山海经图说》和陈连山著作观点，但参考文献列出的分别更像书评/评介文章。只引用评介容易形成二手转引，也不足以支撑异兽考据。",
    "增加实际采用的《山海经》权威校注本、图像学资料及所述专著；凡通过二手文献获知的观点要核对原书。涉及成书年代等存在争议的问题，应呈现不同观点并说明本项目采用哪一版本作为创作依据。",
)
add_rec(
    doc,
    "逐条核验参考文献真实性与格式",
    "任务书主要参考文献[1]—[16]。",
    "不同条目在作者、空格、标点、文献类型和DOI呈现上不完全统一；正文中“吕析影”与参考文献“吕栎影”写法不一致。",
    "以知网/期刊官网的题录页逐条核验作者、题名、刊名、年卷期、页码和DOI；统一按学院文件或GB/T 7714格式著录。正文引用应与文后序号一一对应，不要仅在综述中提作者而无引文序号。",
    None,
    "P0",
)
add_rec(
    doc,
    "让“42篇、27部、11部”可复核",
    "1.3.1与1.3.2中的数量表述。",
    "数量看似具体，但缺少检索式、时间、纳入标准和样本表，答辩时容易被追问。",
    "把检索记录和作品样本清单放入附录。若无法提供完整证据，应改成较谨慎表述，或只报告已实际完成编码分析的样本数。",
)

doc.add_heading("五、语言、格式与文档状态建议", level=1)
add_rec(
    doc,
    "统一学术语体和术语",
    "全文。",
    "存在“重中之重”“有很大帮助”“因为IP价值浓厚”“浏览……讨论得出”等口语或宣传化表达；AI短剧、AI微短剧、AIGC短剧混用。",
    "统一核心术语并在绪论首次出现时界定；改用“本项目拟……”“分析表明……”“用于支撑……”等可验证表达。删减“完全、所有、严格、最终”等难以证明的绝对词。",
)
add_rec(
    doc,
    "清理修订痕迹、空段和个人元数据",
    "文档底层结构与表格单元格。",
    "结构检查发现文档含46处插入型修订标记；开题报告主体表格中存在大量空段（其中两处主要内容单元格分别有24个和21个空段）；文档作者/最后修改者元数据与学生姓名不一致。",
    "在保留备份的前提下，提交版接受或拒绝全部修订，关闭修订模式；删除人为撑版的空段，使用段前段后和分页设置控制版式；通过“检查文档/检查问题”清理作者、最后修改者和批注等个人信息。",
    None,
    "P0",
)
add_rec(
    doc,
    "检查表格分页和标题层级",
    "任务书大表、开题报告三张单列表格及各级标题。",
    "正文几乎全部放在单列表格内，内容很长；大量空段可能造成异常分页。标题有“1.3资料准备—1.3.1……”与拟定论文目录“1.2文献综述—1.2.1……”两套层级，容易混淆。",
    "保留学校模板时不要改表格框架，但应启用允许跨页断行、合理设置单元格内段落间距，避免整段被推到下一页；统一开题报告内部编号。最终设计阐述的目录编号可另行设计，不必机械复制开题报告编号。",
)

doc.add_heading("六、建议的修改顺序", level=1)
steps = [
    "先向指导教师确认：个人/团队属性、最终成果是20集成片还是策划加样片、单集时长与学校技术规格。",
    "统一任务书与开题报告中的题目、术语、主语、成果形态和时间节点。",
    "补写可验收技术指标、分集矩阵、异兽考据卡和个人分工证明。",
    "重写文献综述：分类归纳、比较观点、提出研究缺口，并核验全部题录。",
    "补充研究方法、AIGC工作流、评价方案与版权/伦理风险控制。",
    "调整设计阐述提纲，使“方法—过程—结果—评价—反思”形成闭环。",
    "最后做语言精修、修订接受、空段清理、元数据清理和版式检查。",
]
for i, text in enumerate(steps, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.32)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"{i}. ")
    set_run_font(r, bold=True, color=BLUE)
    r = p.add_run(text)
    set_run_font(r)

doc.add_heading("七、开题答辩前自查清单", level=1)
checks = [
    "题目、成果形式、集数和单集时长在任务书与开题报告中完全一致。",
    "全文已统一使用“本人/本项目”或明确团队分工，不再混用“我们/小组/我的工作”。",
    "三项核心设计问题均有对应方法、过程材料和评价指标。",
    "20集分集矩阵、异兽考据卡、作品样本表、提示词迭代记录已有附件或样例。",
    "所有统计数量和市场判断均有来源、样本范围或检索说明。",
    "参考文献作者、题名、期刊、年份、卷期、页码、DOI均已逐条核验。",
    "音乐、字体、配音、图像参考和AI工具使用条款均有授权/合规记录。",
    "已接受/拒绝全部修订，删除空段，清除不相关作者元数据，检查分页和表格。",
]
for text in checks:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.30)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("□ ")
    set_run_font(r, bold=True, color=BLUE)
    r = p.add_run(text)
    set_run_font(r)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— 完 —")
set_run_font(r, size=10, color=MUTED)

doc.save(OUT)
print(OUT)
