from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path(r"C:\Users\32933\Documents\ChatGPT\论文\outputs\commitment_letter_20260902")
OUT_PATH = OUT_DIR / "承诺书_中国科协博士生专项计划.docx"


def set_run_font(run, east_asia="仿宋_GB2312", western="Times New Roman", size=16, bold=False):
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), western)


def set_paragraph_font(paragraph, **kwargs):
    for run in paragraph.runs:
        set_run_font(run, **kwargs)


def set_east_asia_style(style, east_asia, western, size):
    style.font.name = western
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), western)


def set_document_grid(section, line_pitch_twips=480):
    sect_pr = section._sectPr
    doc_grid = sect_pr.find(qn("w:docGrid"))
    if doc_grid is None:
        doc_grid = OxmlElement("w:docGrid")
        sect_pr.append(doc_grid)
    doc_grid.set(qn("w:type"), "lines")
    doc_grid.set(qn("w:linePitch"), str(line_pitch_twips))


def add_body_paragraph(doc, text, after_pt=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Pt(32)
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(after_pt)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(31)
    run = p.add_run(text)
    set_run_font(run, east_asia="仿宋_GB2312", western="Times New Roman", size=16)
    return p


def add_spacer(doc, height_pt):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(height_pt)
    r = p.add_run(" ")
    set_run_font(r, size=1)
    return p


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.15)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.65)
    section.right_margin = Cm(2.45)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    set_document_grid(section)

    normal = doc.styles["Normal"]
    set_east_asia_style(normal, "仿宋_GB2312", "Times New Roman", 16)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(31)

    # Centered-title pattern, adapted to the formal Chinese commitment-letter reference.
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(34)
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title.paragraph_format.line_spacing = Pt(34)
    title_run = title.add_run("承诺书")
    set_run_font(title_run, east_asia="黑体", western="Arial", size=22, bold=True)

    first = (
        "本人________，系西安交通大学________专业在读博士研究生，学号________，"
        "身份证号________________。本人已充分阅读并理解《中国科协青年科技人才培育工程博士生专项计划》"
        "的相关要求与规定。因本人目前距离博士毕业时间不满两年，但立志于长期在国内从事科学技术研究与创新工作，"
        "现自愿申请参加本项目，并郑重作出如下承诺："
    )
    add_body_paragraph(doc, first, after_pt=13)

    second = (
        "本人承诺，在获得入选资格后，将严格遵守项目培养周期安排，保证在攻读博士学位期间，"
        "完成项目规定的各项培养任务，积极参加中国科协青年科技人才培育工程博士生专项培养活动。"
    )
    add_body_paragraph(doc, second)

    add_spacer(doc, 78)

    sign = doc.add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sign.paragraph_format.right_indent = Cm(0.65)
    sign.paragraph_format.space_after = Pt(16)
    sign.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    sign.paragraph_format.line_spacing = Pt(28)
    sign_run = sign.add_run("承诺人：________")
    set_run_font(sign_run, east_asia="仿宋_GB2312", western="Times New Roman", size=16)

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date.paragraph_format.right_indent = Cm(0.65)
    date.paragraph_format.space_after = Pt(0)
    date.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    date.paragraph_format.line_spacing = Pt(28)
    date_run = date.add_run("____年____月____日")
    set_run_font(date_run, east_asia="仿宋_GB2312", western="Times New Roman", size=16)

    core = doc.core_properties
    core.title = "承诺书"
    core.subject = "中国科协青年科技人才培育工程博士生专项计划"
    core.author = ""
    core.keywords = "承诺书；博士生专项计划；中国科协"

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
